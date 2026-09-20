"""File ingestion: turn a supported input file into the plain text the pipeline already expects.

Called once per input file by the CLI, before Stage 1. No stage knows or cares what the original format
was: every path ends in the same plain text a .txt file would have given.

Supported formats (by extension, then confirmed by content -- a mismatch fails rather than guessing):
  .txt .md   read as UTF-8 (a leading BOM is dropped, exactly as the CLI always did); otherwise unchanged
  .pdf       text layer extracted with pypdf, then cleaned: repeated page headers/footers and standalone
             page-number lines removed, words hyphenated across line breaks re-joined
  .json      one of the explicit shapes in ``_from_json`` (anything else is an error naming those shapes)
  .docx      paragraphs via python-docx; Heading/Title styles become markdown headings ("#", "##", ...)
             so Stage 1's heading-based segmentation can use the document's real structure
Anything else -- other extensions, no extension, binary content in a text file, a ".pdf" that is not a
PDF -- raises IngestionError with a message saying what is supported.

Why pypdf rather than pdfplumber: pypdf is pure Python (no native build) and already used in this project;
it extracts text in content-stream order, which for most generated reports is the reading order, column
by column. pdfplumber rebuilds lines from x/y positions, which merges side-by-side columns line by line.
Known PDF limitations are listed in README / FINDINGS: no OCR; multi-column pages whose content stream is
not in reading order come out interleaved; tables come out as flattened lines.
"""

from __future__ import annotations

import json
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED = {".txt": "plain text", ".md": "Markdown", ".pdf": "PDF", ".json": "JSON", ".docx": "Word (.docx)"}
MIN_PDF_CHARS = 200            # below this for the whole PDF: treated as having no text layer
MIN_PDF_CHARS_PER_PAGE = 25    # ... or below this per page on average
HEADER_FOOTER_ZONE = 3         # non-empty lines at the top and bottom of each page checked for repetition
HEADER_FOOTER_SHARE = 0.5      # a line repeated on at least this share of pages (and >= 3 pages) is stripped


class IngestionError(ValueError):
    """The file cannot be turned into plain text; the message says why and what is supported."""


@dataclass
class Ingested:
    text: str
    format: str                       # "text", "markdown", "pdf", "json", "docx"
    summary: str                      # one line for --verbose, e.g. "Ingested as PDF: ..."
    details: dict = field(default_factory=dict)


def supported_list() -> str:
    return ", ".join(f"{ext} ({name})" for ext, name in SUPPORTED.items())


def ingest(path: str | Path) -> Ingested:
    """Detect the format of ``path`` and return its plain text. Raises IngestionError on anything unsupported."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        what = f"'{ext}'" if ext else "no file extension"
        raise IngestionError(f"{path.name}: unsupported input format ({what}). Supported formats: {supported_list()}. "
                             f"Convert the file to one of these first.")
    data = path.read_bytes()
    if ext in (".txt", ".md"):
        return _from_text(path, data, ext)
    if ext == ".pdf":
        return _from_pdf(path, data)
    if ext == ".json":
        return _from_json(path, data)
    return _from_docx(path, data)


# ------------------------------------------------------------------------------------------ plain text
def _from_text(path: Path, data: bytes, ext: str) -> Ingested:
    if b"\x00" in data:
        raise IngestionError(f"{path.name}: contains binary data (NUL bytes), so it is not a text file despite its "
                             f"'{ext}' extension. Supported formats: {supported_list()}.")
    try:
        # Same call the CLI always used: utf-8-sig drops a leading BOM (Windows editors add one, and it would
        # otherwise stick to the first heading); newlines are normalised to "\n" by text-mode reading.
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise IngestionError(f"{path.name} is not UTF-8 text ({exc}). Convert it to UTF-8 first.") from exc
    kind = "markdown" if ext == ".md" else "text"
    return Ingested(text, kind, f"Ingested as {SUPPORTED[ext]}: {len(text):,} characters, unchanged",
                    {"characters": len(text)})


# ------------------------------------------------------------------------------------------ PDF
# "42", "Page 42", "page 42 of 187", "12 / 40", and lowercase roman front-matter numbers i..xxxix. Uppercase roman
# numerals are NOT matched: a standalone "IV" is far more often a section heading than a page number, and a
# case-insensitive letter class would also take words such as "Mix" or "Civil".
_PAGE_NUMBER = re.compile(r"^\s*(?:(?:[Pp]age|PAGE)\s+)?(?:\d{1,4}|(?=[ivx])x{0,3}(?:ix|iv|v?i{0,3}))"
                          r"(?:\s*(?:of|/)\s*\d{1,4})?\s*$")
_HYPHEN_BREAK = re.compile(r"([a-z])-\n([a-z])")


HEADER_FOOTER_MAX_WORDS = 8     # a line that differs between pages only in its numbers must be this short ...
_SENTENCE_END = re.compile(r"[.!?]['\")\]]*$")   # ... and must not end like a sentence


def _normalise(line: str) -> str:
    """Key for header/footer matching.

    Lines repeated exactly (ignoring case and spacing) match as they are. A line may also match with its numbers
    ignored ('Annual Report 2026 | 14' ~ '... | 15'), but only when it is short and does not end like a sentence:
    otherwise body lines such as 'Revenue rose 4 percent.' / 'Revenue rose 6 percent.' at the top of consecutive
    pages would be taken for a running header and deleted."""
    exact = re.sub(r"\s+", " ", line.strip().lower())
    if len(exact.split()) <= HEADER_FOOTER_MAX_WORDS and not _SENTENCE_END.search(exact):
        return re.sub(r"\d+", "#", exact)
    return exact


def clean_pdf_pages(pages: list[str]) -> tuple[str, dict]:
    """Strip repeated headers/footers and page-number lines from per-page text; join pages. Pure function."""
    page_lines = [[ln.rstrip() for ln in page.splitlines()] for page in pages]
    zone_counts: Counter = Counter()
    for lines in page_lines:
        non_empty = [ln for ln in lines if ln.strip()]
        zone = non_empty[:HEADER_FOOTER_ZONE] + non_empty[-HEADER_FOOTER_ZONE:]
        zone_counts.update({_normalise(ln) for ln in zone if not _PAGE_NUMBER.match(ln)})
    threshold = max(3, HEADER_FOOTER_SHARE * len(pages))
    repeated = {key for key, n in zone_counts.items() if n >= threshold and key}
    stripped_repeated = stripped_numbers = 0
    kept_pages = []
    for lines in page_lines:
        non_empty_idx = [i for i, ln in enumerate(lines) if ln.strip()]
        zone_idx = set(non_empty_idx[:HEADER_FOOTER_ZONE] + non_empty_idx[-HEADER_FOOTER_ZONE:])
        out = []
        for i, ln in enumerate(lines):
            if ln.strip() and _PAGE_NUMBER.match(ln):
                stripped_numbers += 1
                continue
            if i in zone_idx and _normalise(ln) in repeated:
                stripped_repeated += 1
                continue
            out.append(ln)
        kept_pages.append("\n".join(out).strip())
    text = "\n\n".join(p for p in kept_pages if p)
    rejoined = len(_HYPHEN_BREAK.findall(text))
    text = _HYPHEN_BREAK.sub(r"\1\2", text)
    return text, {"repeated_header_footer_lines": stripped_repeated, "page_number_lines": stripped_numbers,
                  "repeated_patterns": sorted(repeated), "hyphenated_breaks_rejoined": rejoined}


def _from_pdf(path: Path, data: bytes) -> Ingested:
    if not data.startswith(b"%PDF"):
        raise IngestionError(f"{path.name}: does not look like a PDF (no %PDF header). Supported formats: "
                             f"{supported_list()}.")
    from io import BytesIO

    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise IngestionError(f"{path.name}: the PDF is encrypted and cannot be read without a password.") from exc
        pages = [page.extract_text() or "" for page in reader.pages]
    except IngestionError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise IngestionError(f"{path.name}: the PDF could not be parsed ({type(exc).__name__}: {exc}).") from exc
    raw_chars = sum(len(p.strip()) for p in pages)
    if not pages or raw_chars < MIN_PDF_CHARS or raw_chars / len(pages) < MIN_PDF_CHARS_PER_PAGE:
        raise IngestionError(
            f"{path.name}: this PDF has no usable text layer ({raw_chars} characters of text across {len(pages)} "
            f"page(s)) -- it looks scanned or image-based. It would need OCR, which is not supported. Convert it "
            f"to text with an OCR tool first.")
    text, info = clean_pdf_pages(pages)
    summary = (f"Ingested as PDF: extracted {len(text):,} characters from {len(pages)} pages, stripped "
               f"{info['repeated_header_footer_lines']} repeated header/footer lines and {info['page_number_lines']} "
               f"page-number lines, re-joined {info['hyphenated_breaks_rejoined']} hyphenated line breaks")
    return Ingested(text, "pdf", summary, {"pages": len(pages), "characters": len(text), **info})


# ------------------------------------------------------------------------------------------ JSON
JSON_SHAPES = ('{"text": "<document text>"}', '{"document": "<document text>"}',
               '{"sections": ["<section text>", ...]}',
               '{"sections": [{"title": "<optional heading>", "text": "<section text>"}, ...]}')


def _from_json(path: Path, data: bytes) -> Ingested:
    expected = "Expected one of: " + " | ".join(JSON_SHAPES) + " (an optional top-level \"title\" string is allowed)."
    try:
        obj = json.loads(data.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise IngestionError(f"{path.name}: not valid UTF-8 JSON ({exc}). {expected}") from exc
    if not isinstance(obj, dict):
        raise IngestionError(f"{path.name}: the top level is a {type(obj).__name__}, not an object. {expected}")
    present = [k for k in ("text", "document", "sections") if k in obj]
    if len(present) != 1:
        found = ", ".join(sorted(obj)) or "none"
        why = "none of" if not present else f"more than one of ({', '.join(present)})"
        raise IngestionError(f"{path.name}: has {why} the fields \"text\", \"document\", \"sections\" "
                             f"(fields found: {found}). {expected}")
    key = present[0]
    title = obj.get("title")
    if title is not None and not isinstance(title, str):
        raise IngestionError(f"{path.name}: \"title\" must be a string. {expected}")
    head = f"# {title.strip()}\n\n" if title and title.strip() else ""
    if key in ("text", "document"):
        if not isinstance(obj[key], str) or not obj[key].strip():
            raise IngestionError(f"{path.name}: \"{key}\" must be a non-empty string. {expected}")
        text = head + obj[key]
        return Ingested(text, "json", f"Ingested as JSON (\"{key}\" field): {len(text):,} characters",
                        {"shape": key, "characters": len(text)})
    sections = obj["sections"]
    if not isinstance(sections, list) or not sections:
        raise IngestionError(f"{path.name}: \"sections\" must be a non-empty array. {expected}")
    blocks, titled = [], 0
    for i, sec in enumerate(sections):
        if isinstance(sec, str):
            body, sec_title = sec, None
        elif isinstance(sec, dict) and isinstance(sec.get("text"), str):
            body, sec_title = sec["text"], sec.get("title")
            if sec_title is not None and not isinstance(sec_title, str):
                raise IngestionError(f"{path.name}: sections[{i}].title must be a string. {expected}")
        else:
            raise IngestionError(f"{path.name}: sections[{i}] is neither a string nor an object with a string "
                                 f"\"text\" field. {expected}")
        if not body.strip():
            continue
        if sec_title and sec_title.strip():
            titled += 1
            blocks.append(f"## {sec_title.strip()}\n\n{body.strip()}")
        else:
            blocks.append(body.strip())
    if not blocks:
        raise IngestionError(f"{path.name}: every section is empty. {expected}")
    text = head + "\n\n".join(blocks)
    return Ingested(text, "json", f"Ingested as JSON (\"sections\" array): {len(blocks)} sections ({titled} with "
                                  f"titles, kept as headings), {len(text):,} characters",
                    {"shape": "sections", "sections": len(blocks), "titled": titled, "characters": len(text)})


# ------------------------------------------------------------------------------------------ DOCX
def _heading_level(style_name: str) -> int | None:
    name = (style_name or "").strip().lower()
    if name == "title":
        return 1
    m = re.match(r"heading\s*(\d)$", name)
    return min(int(m.group(1)) + 1, 6) if m else None     # Title = "#", Heading 1 = "##", ...


def _from_docx(path: Path, data: bytes) -> Ingested:
    if not data.startswith(b"PK\x03\x04"):
        raise IngestionError(f"{path.name}: does not look like a .docx file (not a ZIP package). Supported formats: "
                             f"{supported_list()}.")
    from io import BytesIO

    try:
        import docx
    except ImportError as exc:
        raise IngestionError("Reading .docx needs the python-docx package: pip install python-docx") from exc
    try:
        document = docx.Document(BytesIO(data))
    except Exception as exc:     # python-docx raises several types for a corrupt or non-Word ZIP
        raise IngestionError(f"{path.name}: the .docx file could not be read ({type(exc).__name__}: {exc}).") from exc
    blocks, headings = [], 0
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        level = _heading_level(para.style.name if para.style is not None else "")
        if level:
            headings += 1
            blocks.append(f"{'#' * level} {text}")
        else:
            blocks.append(text)
    if not blocks:
        raise IngestionError(f"{path.name}: the .docx file contains no text paragraphs.")
    text = "\n\n".join(blocks)
    tables = len(document.tables)
    summary = (f"Ingested as Word (.docx): {len(blocks)} paragraphs ({headings} headings kept as markdown headings), "
               f"{len(text):,} characters" + (f"; {tables} table(s) NOT extracted (table text is skipped)" if tables else ""))
    return Ingested(text, "docx", summary, {"paragraphs": len(blocks), "headings": headings,
                                            "tables_skipped": tables, "characters": len(text)})
